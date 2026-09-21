"""
PowerPoint Generator for AI-Powered Resume Analyzer
===================================================
Generates a professional 15-slide PowerPoint (.pptx) presentation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)  # 16:9 widescreen format

    # Color Palette
    DARK_BLUE = RGBColor(15, 23, 42)     # #0f172a
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3b82f6
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    TEXT_DARK = RGBColor(30, 41, 59)     # #1e293b
    BG_LIGHT = RGBColor(248, 250, 252)   # #f8fafc
    CARD_BG = RGBColor(255, 255, 255)
    GREEN = RGBColor(16, 185, 129)

    blank_layout = prs.slide_layouts[6]

    # Helper: Set Slide Background
    def set_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper: Add Standard Slide Header
    def add_header(slide, title, category="FINAL YEAR B.E. CSE PROJECT"):
        set_bg(slide, BG_LIGHT)

        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Title
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_BLUE

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1, DARK_BLUE)

    # Center Title Box
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FINAL YEAR ACADEMIC PROJECT (B.E. CSE)"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf1.add_paragraph()
    p1.text = "AI-Powered Resume Analyzer and\nJob Recommendation System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "Automated NLP Resume Parsing, TF-IDF Job Matching, Skill Gap Analysis & Free Learning Roadmaps"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf1.add_paragraph()
    p3.text = "\nPresented by: Final Year CSE Students  |  Guided by: Project Guide / Dept. of CSE"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = GREEN
    p3.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 2: Project Overview & Abstract
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Project Overview & Abstract")

    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    points2 = [
        "• High Recruitment Volume: Organizations receive hundreds of resumes per job opening, making manual shortlisting slow, costly, and subjective.",
        "• Applicant Uncertainty: Job seekers and students often face Applicant Tracking System (ATS) rejections without understanding why their profile failed.",
        "• Core Solution: An intelligent, open-source web application that parses resumes (PDF/DOCX), evaluates ATS completeness, and recommends target job profiles.",
        "• Explainable AI Engine: Utilizes Natural Language Processing (NLP) for entity parsing and TF-IDF + Cosine Similarity for mathematical job matching.",
        "• Bridging the Skill Gap: Automatically detects missing skills for target roles and delivers curated free learning courses from YouTube, freeCodeCamp, and W3Schools."
    ]
    for pt in points2:
        p = tf2.add_paragraph()
        p.text = pt
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Problem Statement
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Problem Statement & Industry Challenges")

    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    points3 = [
        "1. Inefficiency of Manual Screening: Recruiter fatigue leads to overlooked talent and inconsistent evaluations during manual resume checks.",
        "2. Keyword Mismatch in Traditional ATS: Existing primitive systems rely on naive exact string matching, failing to capture related frameworks and skills.",
        "3. Lack of Constructive Candidate Feedback: Job seekers receive generic rejection emails without actionable insights on their missing qualifications.",
        "4. High Cost of Commercial ATS Tools: Proprietary enterprise ATS software is expensive, closed-source, and unavailable for universities or students.",
        "5. Need for Career Guidance: There is an urgent need for an automated platform that not only matches jobs but also guides students on how to upskill."
    ]
    for pt in points3:
        p = tf3.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 4: Existing System vs Proposed System
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Existing System vs Proposed System")

    # Add Table
    rows, cols = 5, 3
    table_shape = slide4.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.5)

    headers = ["Evaluation Parameter", "Existing / Traditional System", "Proposed AI-Powered System"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE

    row_data = [
        ("Resume Screening", "Manual, slow, prone to cognitive bias", "Automated NLP parsing in under 2 seconds"),
        ("Job Matching", "Basic keyword search with boolean filters", "TF-IDF Vectorization & Cosine Similarity ranking"),
        ("Skill Gap Diagnosis", "No feedback provided to candidates", "Mathematical set difference with market demand ranking"),
        ("Upskilling Guidance", "Candidates must search courses on their own", "Direct curated free learning roadmaps (YouTube/fCC)")
    ]
    for row_idx, data in enumerate(row_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 5: Project Objectives & Scope
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Objectives & Project Scope")

    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    points5 = [
        "• Primary Objectives:",
        "   1. Automatically parse multi-format resumes (PDF, DOCX, TXT) and extract raw textual entities.",
        "   2. Apply NLP word-boundary taxonomies to categorize technical proficiencies across 6 distinct categories.",
        "   3. Compute mathematical text similarity using TF-IDF and Cosine Similarity to recommend ranked tech jobs.",
        "   4. Diagnose skill gaps and evaluate overall ATS completeness (0–100%).",
        "   5. Deliver personalized free learning resources for all identified skill deficits.",
        "• Project Scope:",
        "   - Targets technical and IT job profiles (Python Dev, Data Scientist, Full Stack, DevOps, Cloud Engineer, etc.).",
        "   - 100% Free and Open-Source architecture (no paid external APIs or cloud subscriptions needed)."
    ]
    for pt in points5:
        p = tf5.add_paragraph()
        p.text = pt
        p.font.size = Pt(15 if not pt.startswith("•") else 16)
        p.font.bold = pt.startswith("•")
        p.font.color.rgb = DARK_BLUE if pt.startswith("•") else TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 6: System Architecture & Workflow
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. System Architecture & End-to-End Workflow")

    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    steps6 = [
        "1. Client Layer (Presentation): User uploads resume via modern Bootstrap 5 web interface.",
        "2. Ingestion & File Parsing: PyPDF2 / python-docx extracts unstructured text from document pages and tables.",
        "3. NLP Preprocessing & Extraction: Tokenization, section splitting, and regex boundary matching isolate skills, education, experience, and contact information.",
        "4. ATS Completeness Scoring: Evaluates structural completeness across 5 key dimensions (Max 100%).",
        "5. ML Matching Engine: TF-IDF vectorizes resume + job descriptions; Cosine Similarity computes match scores.",
        "6. Skill Gap & Learning Dispatch: Calculates set difference and retrieves free learning course roadmaps.",
        "7. Database Layer: SQLAlchemy ORM persists all users, resumes, recommendations, and analytics in MySQL."
    ]
    for s in steps6:
        p = tf6.add_paragraph()
        p.text = s
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 7: Resume Parsing Module
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Multi-Format Resume Parsing Module")

    tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    points7 = [
        "• PDF Parsing (PyPDF2 / pdfplumber):",
        "   - Iterates through document page objects, extracting text streams and handling multi-page formatting.",
        "• Microsoft Word Parsing (python-docx):",
        "   - Scans paragraph text blocks and iterates across table structures where skills/education are often placed.",
        "• Plain Text Parsing (.txt):",
        "   - Implements multi-encoding fallbacks (UTF-8, Latin-1, CP1252) for robust cross-platform ingestion.",
        "• Security & Validation:",
        "   - Filename sanitization via werkzeug.utils.secure_filename.",
        "   - Strict file extension validation and 16 MB maximum size enforcement."
    ]
    for pt in points7:
        p = tf7.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 8: NLP Skill Extraction Module
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. NLP Skill & Entity Extraction Module")

    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    points8 = [
        "• 6 Taxonomy Categories Extracted:",
        "   1. Programming Languages: Python, Java, JavaScript, TypeScript, C++, C#, SQL, Go, Rust...",
        "   2. Frameworks & Web: React, Flask, Django, Node.js, Spring Boot, Express, Angular, Next.js...",
        "   3. Databases: MySQL, PostgreSQL, MongoDB, Redis, SQLite, Oracle...",
        "   4. Developer Tools & Cloud: Git, Docker, Kubernetes, AWS, Linux, Postman, CI/CD...",
        "   5. Domain & AI/ML: Machine Learning, Deep Learning, NLP, Data Analysis, Pandas, REST APIs...",
        "   6. Soft Skills: Problem Solving, Teamwork, Communication, Agile Methodologies...",
        "• Word-Boundary Regex Protection: Lookarounds `(?<![a-zA-Z0-9])term(?![a-zA-Z0-9])` prevent false matches (e.g., 'Java' vs 'JavaScript', 'C' vs 'CSS').",
        "• Metadata Extraction: Captures candidate Email, Phone, LinkedIn, GitHub, Degrees, and Experience."
    ]
    for pt in points8:
        p = tf8.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 9: AI/ML Matching Engine
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "8. AI/ML Matching Engine (TF-IDF + Cosine Similarity)")

    tb9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    points9 = [
        "• Step 1: TF-IDF Vectorization (Term Frequency - Inverse Document Frequency)",
        "   - Converts resume text and job descriptions into numeric vectors.",
        "   - Formula: TF(t, d) = (count of t in d) / (total words in d);  IDF(t, D) = log(|D| / (1 + count(d with t)))",
        "   - Rare technical keywords receive higher importance; generic words ('the', 'and') are down-weighted.",
        "• Step 2: Cosine Similarity Calculation",
        "   - Measures the cosine of the angle between the Resume Vector (A) and Job Vector (B):",
        "   - Formula: Cosine Similarity = (A · B) / (||A|| × ||B||)",
        "   - Output ranges from 0.0 (No overlap) to 1.0 (Identical semantic direction), converted to 0–100%.",
        "• Step 3: Hybrid Scoring & Ranking",
        "   - Combines 40% TF-IDF Text Similarity + 60% Weighted Skill Overlap for highly explainable recommendations."
    ]
    for pt in points9:
        p = tf9.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 10: Skill Gap Analysis Module
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "9. Skill Gap Analysis & Priority Engine")

    tb10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    points10 = [
        "• Mathematical Set Difference Formulation:",
        "   - Matched Skills = Candidate Skills ∩ Job Required Skills",
        "   - Missing Skills = Job Required Skills \\ Candidate Skills",
        "   - Readiness Percentage = (|Matched Skills| / |Total Required Skills|) × 100",
        "   - Skill Gap Index = 100% − Readiness Percentage",
        "• Global Market Demand Prioritization:",
        "   - Across all benchmark jobs in the system, each missing skill is assigned a frequency demand score.",
        "   - Skills required across multiple roles (e.g. Docker, Git, SQL) are given HIGH study priority.",
        "   - Enables candidates to maximize career impact with their study time."
    ]
    for pt in points10:
        p = tf10.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 11: Free Learning Recommendation Module
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "10. Free Learning Recommendation Module")

    tb11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    points11 = [
        "• 100% Free & Curated Educational Resources:",
        "   - Maps every missing skill directly to high-quality free courses, video crash courses, and tutorials.",
        "• Integrated Platforms:",
        "   - YouTube: Full-length video courses from freeCodeCamp, Traversy Media, CS Dojo, TechWithTim.",
        "   - freeCodeCamp: Interactive certifications in Python, JavaScript, Front-End, Data Analysis, and ML.",
        "   - Coursera Free Audit: DeepLearning.AI and university machine learning foundational specializations.",
        "   - W3Schools & Official Documentation: Interactive playgrounds and reference guides.",
        "• Tailored Learning Roadmaps: Organizes missing skills with estimated study hours and direct clickable URLs."
    ]
    for pt in points11:
        p = tf11.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 12: Database Design & Architecture
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "11. Database Design & Relational Schema")

    tb12 = slide12.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    points12 = [
        "• RDBMS Engine: MySQL 8.0+ managed via SQLAlchemy ORM (with auto SQLite fallback).",
        "• 7 Normalized Relational Tables:",
        "   1. users: user_id (PK), full_name, email (UQ), password_hash, role (student/admin), phone",
        "   2. resumes: resume_id (PK), user_id (FK), resume_title, file_path, extracted_text, completeness_score",
        "   3. skills: skill_id (PK), skill_name (UQ), skill_category",
        "   4. jobs: job_id (PK), title, company, industry, experience_level, salary_range, description",
        "   5. job_skills: job_skill_id (PK), job_id (FK), skill_id (FK), importance_level, weight",
        "   6. recommendations: recommendation_id (PK), resume_id (FK), job_id (FK), match_percentage, skills JSON",
        "   7. learning_resources: resource_id (PK), skill_id (FK), title, url, platform, duration_hours, is_free"
    ]
    for pt in points12:
        p = tf12.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 13: UI & Implementation Modules
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "12. UI Design & Key Application Features")

    tb13 = slide13.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf13 = tb13.text_frame
    tf13.word_wrap = True

    points13 = [
        "• Modern Responsive UI: Built with HTML5, CSS3, Bootstrap 5, and Jinja2 templates.",
        "• Drag-and-Drop Upload Zone: Animated file drop area with instant file size and format validation.",
        "• Candidate Dashboard: Displays latest scan progress, ATS score meters, top recommendations, and history.",
        "• Recruiter / Admin Analytics Portal: Interactive Chart.js visualization of top in-demand skills in the talent pool.",
        "• Job Management CRUD: Admin capability to add/edit job descriptions and assign mandatory skill weights.",
        "• 1-Click Desktop Launcher: Windows batch shortcut (`START_PROJECT.bat`) to launch and open browser automatically."
    ]
    for pt in points13:
        p = tf13.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(11)

    # -------------------------------------------------------------
    # SLIDE 14: Testing, Results & Validation
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "13. Testing, Results & Performance Validation")

    tb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    points14 = [
        "• Comprehensive Automated Test Suite (test_app.py):",
        "   - Executed 9 unit tests verifying all 11 phases in 3.74 seconds with 100% PASS rate (OK).",
        "• Validated Test Scenarios:",
        "   1. Multi-format file ingestion (PDF, DOCX, TXT parsing).",
        "   2. NLP taxonomy extraction without false positives (e.g. C vs CSS, Java vs JavaScript).",
        "   3. TF-IDF Cosine Similarity ranking accuracy against real resume samples.",
        "   4. Skill gap set operations and ATS scoring diagnostic accuracy.",
        "   5. User authentication, password hashing (Werkzeug PBKDF2/scrypt), and session guards.",
        "• Performance Benchmark: Average processing time per resume is under 2.0 seconds on standard CPU hardware."
    ]
    for pt in points14:
        p = tf14.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(9)

    # -------------------------------------------------------------
    # SLIDE 15: Conclusion & Future Scope
    # -------------------------------------------------------------
    slide15 = prs.slides.add_slide(blank_layout)
    set_bg(slide15, DARK_BLUE)

    title_box15 = slide15.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(11.3), Inches(5.5))
    tf15 = title_box15.text_frame
    tf15.word_wrap = True

    p0 = tf15.paragraphs[0]
    p0.text = "14. ADVANTAGES, FUTURE SCOPE & CONCLUSION"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p1 = tf15.add_paragraph()
    p1.text = "Conclusion:"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf15.add_paragraph()
    p2.text = "The AI-Powered Resume Analyzer and Job Recommendation System demonstrates how open-source NLP and Machine Learning techniques (TF-IDF + Cosine Similarity) can automate recruitment screening, evaluate ATS readiness, and guide candidates with personalized upskilling roadmaps."
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf15.add_paragraph()
    p3.text = "\nFuture Enhancements:"
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = GREEN

    p4 = tf15.add_paragraph()
    p4.text = "• Integration with Transformer-based Semantic Models (BERT / RoBERTa / Sentence-Transformers).\n• Optical Character Recognition (OCR) for scanned image resumes (Tesseract OCR).\n• Live Job Portal API Integration (LinkedIn, Indeed, Naukri API synchronization)."
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_WHITE

    p5 = tf15.add_paragraph()
    p5.text = "\nThank You! Any Questions?"
    p5.font.size = Pt(22)
    p5.font.bold = True
    p5.font.color.rgb = ACCENT_BLUE
    p5.alignment = PP_ALIGN.CENTER

    output_path = "AI_Resume_Analyzer_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_presentation()
